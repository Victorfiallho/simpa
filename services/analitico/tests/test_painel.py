"""Painel, indicadores, relatórios e histórico de risco (D11), com o acadêmico simulado."""

import io
from decimal import Decimal as D

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.clientes import academico
from app.core.database import Base, get_db
from app.core.seguranca import UsuarioAtual, usuario_atual
from app.main import app
from app.models import ResultadoRisco
from app.services import risco

TURMA = {"id": 1, "codigo": "IA2-2026.2-A", "nome": "IA · 2º período A", "status": "ativa"}


def inscricao(id_, aluno_id, nome, sigla, notas=(), aulas=32, faltas=0, status="ativa"):
    """Monta uma inscrição como a rota /internal do acadêmico devolveria."""
    return academico.Inscricao.model_validate(
        {
            "id": id_,
            "status": status,
            "aluno": {
                "id": aluno_id,
                "matricula": f"2026{aluno_id:04d}",
                "nome_completo": nome,
                "email": f"aluno{aluno_id}@aluno.simpa.dev",
                "turma_id": 1,
                "status": "ativo",
            },
            "disciplina": {
                "id": {"EST": 1, "POO": 2, "ALG": 3}[sigla],
                "codigo": f"{sigla}-2026.2-A",
                "nome": {"EST": "Estatística", "POO": "POO", "ALG": "Álgebra"}[sigla],
                "periodo_letivo": "2026.2",
                "turma_id": 1,
                "professor_id": 3,
                "status": "ativa",
                "turma": TURMA,
            },
            "avaliacoes": [
                {"periodo_avaliacao": p, "nota": n, "peso": peso} for p, n, peso in notas
            ],
            "frequencias": [{"periodo_avaliacao": "N1", "aulas_previstas": aulas, "faltas": faltas}]
            if notas
            else [],
        }
    )


BOM = [("N1", "8.00", "4"), ("N2", "9.00", "6")]
INSCRICOES = [
    inscricao(1, 1, "Ana Lima", "EST", BOM),
    inscricao(2, 1, "Ana Lima", "ALG", [("N1", "4.00", "4"), ("N2", "5.00", "6")]),
    inscricao(3, 2, "Bruno Costa", "EST", [("N1", "6.00", "4"), ("N2", "6.00", "6")], faltas=8),
    inscricao(4, 2, "Bruno Costa", "ALG", BOM),
    inscricao(5, 3, "Clara Souza", "EST"),  # sem dados
    inscricao(6, 3, "Clara Souza", "ALG", BOM, status="cancelada"),  # não conta
]


@pytest.fixture
def cliente(monkeypatch):
    monkeypatch.setattr(academico, "buscar_inscricoes", lambda **_: INSCRICOES)
    app.dependency_overrides[usuario_atual] = lambda: UsuarioAtual(
        id=1, role="coordenador", nome="coordenacao"
    )
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_painel_consolida_alunos_e_situacoes(cliente):
    painel = cliente.get("/painel").json()["dados"]

    assert painel["total_alunos"] == 3
    assert painel["total_inscricoes"] == 5  # a cancelada fica de fora
    assert painel["inscricoes_sem_dados"] == 1
    assert painel["contagem"] == {
        "em_risco": 1,
        "atencao": 1,
        "sem_risco": 0,
        "indeterminado": 1,
    }
    ana = next(a for a in painel["alunos"] if a["nome"] == "Ana Lima")
    assert ana["situacao"] == "em_risco"
    assert "Álgebra: média 4,60 (< 6,00)" in ana["inscricoes"][1]["justificativa"]


def test_painel_traz_dados_dos_graficos(cliente):
    painel = cliente.get("/painel").json()["dados"]

    assert sum(f["inscricoes"] for f in painel["histograma"]) == 4
    assert len(painel["dispersao"]) == 4
    assert {o["oferta"] for o in painel["situacao_por_oferta"]} == {"EST A", "ALG A"}
    # Médias válidas: 8,60 (EST Ana), 4,60 (ALG Ana), 6,00 (EST Bruno), 8,60 (ALG Bruno)
    assert painel["media_geral"] == "6.95"  # 27,80 / 4


def test_indicadores_da_turma(cliente):
    dados = cliente.get("/indicadores/turmas/1").json()["dados"]

    assert dados["inscricoes_validas"] == 4
    assert dados["media"]["tipo_variancia"] == "populacional"
    assert dados["media"]["n"] == 4
    assert dados["por_disciplina"][0]["media"] >= dados["por_disciplina"][-1]["media"]


def test_janela_parcial_usa_variancia_amostral(cliente):
    dados = cliente.get("/indicadores/turmas/1", params={"janela": "N1"}).json()["dados"]

    assert dados["janela"] == "N1"
    assert dados["media"]["tipo_variancia"] == "amostral"


def test_sem_token_responde_401():
    resposta = TestClient(app).get("/painel")

    assert resposta.status_code == 401
    assert resposta.json()["erro"]["codigo"] == "NAO_AUTENTICADO"


def test_relatorio_csv_tem_cabecalho_e_email_mascarado(cliente):
    resposta = cliente.get("/relatorios/inscricoes", params={"formato": "csv"})

    texto = resposta.content.decode("utf-8-sig")
    assert resposta.headers["content-type"].startswith("text/csv")
    assert "attachment" in resposta.headers["content-disposition"]
    assert "Gerado por;coordenacao (coordenador)" in texto
    assert "al***@aluno.simpa.dev" in texto
    assert "aluno1@" not in texto
    assert "4,60" in texto  # vírgula decimal


def test_relatorio_filtra_por_situacao(cliente):
    texto = cliente.get(
        "/relatorios/inscricoes", params={"formato": "csv", "situacao": "em_risco"}
    ).content.decode("utf-8-sig")

    linhas_de_dados = [linha for linha in texto.splitlines() if "@aluno.simpa.dev" in linha]
    assert len(linhas_de_dados) == 1
    assert "Álgebra" in linhas_de_dados[0]


def test_relatorio_xlsx(cliente):
    resposta = cliente.get("/relatorios/inscricoes", params={"formato": "xlsx"})

    planilha = load_workbook(io.BytesIO(resposta.content)).active
    valores = [c.value for linha in planilha.iter_rows() for c in linha if c.value is not None]
    assert "SIMPA — Relatório acadêmico" in valores
    assert "Matrícula" in valores


def test_relatorio_pdf(cliente):
    resposta = cliente.get("/relatorios/inscricoes", params={"formato": "pdf"})

    assert resposta.headers["content-type"] == "application/pdf"
    assert resposta.content.startswith(b"%PDF")


# ---------- Histórico de risco (D11) ----------


@pytest.fixture
def sessao():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    fabrica = sessionmaker(bind=engine, expire_on_commit=False)
    with fabrica() as db:
        yield db
    engine.dispose()


def test_so_grava_quando_a_classificacao_muda(sessao):
    primeira = risco.registrar(sessao, INSCRICOES[0])
    repetida = risco.registrar(sessao, INSCRICOES[0])
    piorou = risco.registrar(
        sessao, inscricao(1, 1, "Ana Lima", "EST", [("N1", "3.00", "4"), ("N2", "4.00", "6")])
    )

    assert primeira is not None
    assert repetida is None
    assert piorou.status == "em_risco"
    assert piorou.criterios_violados == "media_baixa"
    assert piorou.media_final == D("3.60")
    assert sessao.query(ResultadoRisco).count() == 2


def test_inscricao_cancelada_nao_gera_resultado(sessao):
    assert risco.registrar(sessao, INSCRICOES[5]) is None


def test_rota_de_historico(cliente, sessao):
    risco.registrar(sessao, INSCRICOES[1])
    app.dependency_overrides[get_db] = lambda: sessao

    linhas = cliente.get("/riscos/inscricoes/2/historico").json()["dados"]

    assert len(linhas) == 1
    assert linhas[0]["status"] == "em_risco"
    assert linhas[0]["media_final"] == "4.60"
