"""Relatórios de inscrições em CSV, XLSX e PDF (RF11, RNF04).

Cabeçalho com data, hora e quem gerou. O e-mail sai mascarado e o CPF não sai.
"""

import csv
import io
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.seguranca import UsuarioAtual
from app.schemas.analitico import Painel

FUSO = ZoneInfo("America/Sao_Paulo")
ROTULO_SITUACAO = {
    "em_risco": "Em risco",
    "atencao": "Atenção",
    "sem_risco": "Sem risco",
    "indeterminado": "Indeterminado",
}
COLUNAS = [
    "Matrícula",
    "Aluno",
    "E-mail",
    "Turma",
    "Disciplina",
    "N1",
    "N2",
    "Média",
    "Frequência (%)",
    "Situação",
    "Justificativa",
]


def mascarar_email(email: str) -> str:
    """ana.lima@aluno.simpa.dev → an***@aluno.simpa.dev"""
    usuario, _, dominio = email.partition("@")
    return f"{usuario[:2]}***@{dominio}"


@dataclass(frozen=True)
class Relatorio:
    gerado_em: datetime
    gerado_por: str
    filtros: str
    linhas: list[list]


def montar(
    painel: Painel,
    emails: dict[int, str],
    usuario: UsuarioAtual,
    filtros: str,
    situacao: str | None,
) -> Relatorio:
    linhas = []
    for aluno in painel.alunos:
        for inscricao in aluno.inscricoes:
            if situacao and inscricao.situacao != situacao:
                continue
            linhas.append(
                [
                    aluno.matricula,
                    aluno.nome,
                    mascarar_email(emails.get(aluno.id, "")),
                    aluno.turma_nome,
                    inscricao.disciplina_nome,
                    inscricao.nota_n1,
                    inscricao.nota_n2,
                    inscricao.media,
                    inscricao.frequencia,
                    ROTULO_SITUACAO[inscricao.situacao],
                    inscricao.justificativa,
                ]
            )
    quem = f"{usuario.nome or f'usuário {usuario.id}'} ({usuario.role})"
    return Relatorio(datetime.now(FUSO), quem, filtros or "nenhum", linhas)


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, Decimal):
        return f"{valor:.2f}".replace(".", ",")
    return str(valor)


def _cabecalho(relatorio: Relatorio) -> list[list[str]]:
    return [
        ["SIMPA — Relatório acadêmico"],
        ["Gerado em", relatorio.gerado_em.strftime("%d/%m/%Y %H:%M (horário de Brasília)")],
        ["Gerado por", relatorio.gerado_por],
        ["Filtros", relatorio.filtros],
        ["Dados sintéticos de demonstração; e-mails mascarados (LGPD)."],
        [],
    ]


def gerar_csv(relatorio: Relatorio) -> bytes:
    """CSV com ";" e vírgula decimal, que o Excel em português abre direto."""
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";")
    escritor.writerows(_cabecalho(relatorio))
    escritor.writerow(COLUNAS)
    escritor.writerows([[_texto(v) for v in linha] for linha in relatorio.linhas])
    return ("﻿" + saida.getvalue()).encode("utf-8")


def gerar_xlsx(relatorio: Relatorio) -> bytes:
    planilha = Workbook()
    aba = planilha.active
    aba.title = "Inscrições"
    for linha in _cabecalho(relatorio):
        aba.append(linha)
    aba["A1"].font = Font(bold=True, size=14)

    aba.append(COLUNAS)
    linha_titulos = aba.max_row
    for celula in aba[linha_titulos]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="2C2C2C")

    for linha in relatorio.linhas:
        aba.append(linha)
        for celula in aba[aba.max_row][5:9]:
            celula.number_format = "0.00"
        aba.cell(aba.max_row, len(COLUNAS)).alignment = Alignment(wrap_text=True)

    larguras = [12, 26, 26, 22, 30, 7, 7, 8, 14, 13, 70]
    for indice, largura in enumerate(larguras, start=1):
        aba.column_dimensions[aba.cell(1, indice).column_letter].width = largura
    aba.freeze_panes = aba.cell(linha_titulos + 1, 1)

    saida = io.BytesIO()
    planilha.save(saida)
    return saida.getvalue()


def gerar_pdf(relatorio: Relatorio) -> bytes:
    saida = io.BytesIO()
    documento = SimpleDocTemplate(
        saida,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="SIMPA — Relatório acadêmico",
    )
    estilos = getSampleStyleSheet()
    pequeno = estilos["BodyText"].clone("pequeno", fontSize=7, leading=8.5)

    elementos = [
        Paragraph("SIMPA — Relatório acadêmico", estilos["Title"]),
        Paragraph(
            f"Gerado em {relatorio.gerado_em:%d/%m/%Y %H:%M} (horário de Brasília) por "
            f"{relatorio.gerado_por}. Filtros: {relatorio.filtros}. "
            "Dados sintéticos de demonstração; e-mails mascarados.",
            estilos["BodyText"],
        ),
        Spacer(1, 6 * mm),
    ]

    colunas_pdf = [0, 1, 3, 4, 5, 6, 7, 8, 9, 10]  # sem o e-mail, para caber na página
    dados = [[COLUNAS[i] for i in colunas_pdf]]
    for linha in relatorio.linhas:
        dados.append(
            [
                Paragraph(_texto(linha[i]), pequeno) if i in (1, 3, 4, 10) else _texto(linha[i])
                for i in colunas_pdf
            ]
        )

    tabela = Table(
        dados,
        colWidths=[
            20 * mm,
            38 * mm,
            30 * mm,
            38 * mm,
            11 * mm,
            11 * mm,
            12 * mm,
            18 * mm,
            18 * mm,
            77 * mm,
        ],
        repeatRows=1,
    )
    tabela.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2C2C2C")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2EFE7")]),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2DDD0")),
            ]
        )
    )
    elementos.append(tabela)
    documento.build(elementos)
    return saida.getvalue()


GERADORES = {
    "csv": (gerar_csv, "text/csv; charset=utf-8"),
    "xlsx": (gerar_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf": (gerar_pdf, "application/pdf"),
}
